#!/usr/bin/env python3
# ═══════════════════════════════════════════════════════════════
# src/utils/file_text.py — Extracción de texto de archivos (2026-09-01)
# ═══════════════════════════════════════════════════════════════
# Ruta: /ruta/a/axioma/src/utils/file_text.py
# Autor: SaintWick
#
# Propósito: extraer TEXTO REAL de archivos subidos para que los modelos
# los analicen. Antes, los PDF/DOCX se leían como bytes/texto crudo
# (basura binaria o crash de decodificación). Ahora:
#   - .pdf  → pypdf (extracción de texto por página)
#   - .docx → python-docx (párrafos + tablas)
#   - .txt/.md/.py/.json/.csv/.log/.toon/.jsonl → texto utf-8 (con fallback)
# Degradación silenciosa: si falta la librería o falla el parseo, cae a
# texto crudo decodificado con errors='ignore' — nunca lanza hacia el caller.
# ═══════════════════════════════════════════════════════════════
from __future__ import annotations

from pathlib import Path
from typing import Optional

# Límite por defecto para no saturar el contexto del LLM.
DEFAULT_MAX_CHARS = 150_000

# Extensiones que se parsean como documento estructurado.
_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx", ".doc"}


def extract_file_text(
    path: Path,
    max_chars: int = DEFAULT_MAX_CHARS,
    fallback_encoding: str = "utf-8",
) -> str:
    """Extrae texto real de un archivo según su extensión.

    Args:
        path: Ruta del archivo.
        max_chars: Límite de caracteres a devolver (evita saturar el LLM).
        fallback_encoding: Encoding para archivos de texto plano.

    Returns:
        Texto extraído (truncado a max_chars) o "" si el archivo no existe
        / no se puede leer. Nunca lanza.
    """
    if not path or not path.exists():
        return ""

    ext = path.suffix.lower()

    try:
        if ext in _PDF_EXTS:
            text = _extract_pdf(path)
        elif ext in _DOCX_EXTS:
            text = _extract_docx(path)
        else:
            text = _extract_text(path, fallback_encoding)

        if not text:
            # Último recurso: leer binario como texto con errors='ignore'
            text = _read_raw(path)
    except Exception:
        # Degradación total: nunca romper el análisis por un parseo fallido
        text = _read_raw(path)

    if max_chars and max_chars > 0 and len(text) > max_chars:
        text = text[:max_chars] + (
            f"\n\n⚠️ [ARCHIVO TRUNCADO: se muestran {max_chars} de "
            f"~{len(text)} caracteres. Para archivos más grandes, "
            "especificá qué sección querés analizar.]"
        )
    return text


def _extract_pdf(path: Path) -> str:
    """Extrae texto de un PDF con pypdf (degradación silenciosa)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                pages.append(t.strip())
        if not pages:
            return ""
        return "\n\n".join(pages)
    except Exception:
        return ""


def _extract_docx(path: Path) -> str:
    """Extrae texto de un .docx con python-docx (párrafos + tablas)."""
    try:
        import docx  # python-docx
    except ImportError:
        return ""
    try:
        document = docx.Document(str(path))
        parts: list[str] = []
        for para in document.paragraphs:
            if para.text and para.text.strip():
                parts.append(para.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _extract_text(path: Path, encoding: str) -> str:
    """Lee un archivo de texto plano con el encoding pedido."""
    try:
        return path.read_text(encoding=encoding, errors="ignore")
    except Exception:
        return ""


def _read_raw(path: Path) -> str:
    """Último recurso: lee bytes y los decodifica ignorando errores."""
    try:
        data = path.read_bytes()
        return data.decode("utf-8", errors="ignore")
    except Exception:
        return ""
